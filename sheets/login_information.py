"""Login Information -- a simple flat reference table (PLAN.md §14), the
lightest sheet in the workbook by design: no hierarchy, no archiving, no
formulas. A genuine auto-expanding Excel Table, same as Quick Links -- the
user just types directly into most columns, no popup mechanism involved.

**Corrected from §14's original draft (user-requested, PLAN.md §19 Phase
5)**: Notes was originally plain free text, unlike every other sheet's
Notes column. Now reuses the shared `frmNotes` popup exactly as Tasks/
Archive/Campaigns/Campaigns Archive/Recurring Tasks do -- `frmNotes` is
already sheet-agnostic (keyed by sheet name + cell address, not a stable
row key, per `sheets/forms/notes_form.py`), so this needed no new
mechanism, just the same `notes_icon_format()` masking format plus a new
`vba/Login Information.bas` sheet code-behind routing a Notes-column
double-click to it (mirrors `Archive.bas`, the simplest existing example
of a flat sheet with no other interactive columns). Direct typing into
the cell still works too, same as every other Notes column in the
workbook -- only Archive additionally blocks direct edits, and that's for
reasons unrelated to Notes (the whole sheet is a protected historical
log).

**Corrected from §14's original draft, per §17's later/authoritative
treatment**: Access is documented there as "Yes/No, color-coded
(green/red)", but design_system.py's ACCESS_YES/ACCESS_NO (§17.4) already
settled on good-green/muted-gray specifically because "no access" isn't
alarming, just informational -- not literal red. Followed §17 here, since
later sections correct earlier ones per this project's own convention.

Row allowance (CF/DV pre-applied over a fixed range, same reasoning as
every prior sheet -- these are discrete objects that don't travel with
native Table growth): 50 rows, confirmed with the user.
"""

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table

from . import conditional_formatting as cf
from . import design_system as ds

SHEET_NAME = "Login Information"
HEADER_ROW = 1
FIRST_DATA_ROW = 2
LAST_ROW = FIRST_DATA_ROW + 50 - 1  # 51

_DATE_FORMAT = "DD.MM.YYYY"

HEADERS = ["Platform", "Website", "Username / Email", "Password", "Last Password Change", "Access", "Notes"]
_COL = {name: i + 1 for i, name in enumerate(HEADERS)}


def build(wb):
    """Create the Login Information sheet in `wb` (an openpyxl Workbook).
    Returns the created worksheet.
    """
    ws = wb.create_sheet(SHEET_NAME)
    ds.apply_modern_chrome(ws)
    ws.freeze_panes = ws.cell(row=FIRST_DATA_ROW, column=1)
    ws.row_dimensions[HEADER_ROW].height = ds.HEADER_ROW_HEIGHT

    for col, label in enumerate(HEADERS, start=1):
        cell = ws.cell(row=HEADER_ROW, column=col, value=label)
        ds.style_header_cell(cell)

    widths = {
        "Platform": 16,
        "Website": 26,
        "Username / Email": 22,
        "Password": 18,
        "Last Password Change": 16,
        "Access": 10,
        "Notes": 24,
    }
    for name, width in widths.items():
        ws.column_dimensions[get_column_letter(_COL[name])].width = width

    ws.cell(row=FIRST_DATA_ROW, column=_COL["Last Password Change"]).number_format = _DATE_FORMAT

    # Notes: shared frmNotes popup (see module docstring's correction
    # note) -- masking format applied across the full 50-row allowance
    # upfront, same as this sheet's own DV/CF below, since there's no
    # hierarchy engine here to apply it row-by-row as rows are created.
    notes_format = ds.notes_icon_format()
    for row in range(FIRST_DATA_ROW, LAST_ROW + 1):
        ws.cell(row=row, column=_COL["Notes"]).number_format = notes_format

    access_letter = get_column_letter(_COL["Access"])
    access_dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    access_dv.add(f"{access_letter}{FIRST_DATA_ROW}:{access_letter}{LAST_ROW}")
    ws.add_data_validation(access_dv)

    cf.apply_rules(ws, cf.binary_state_rules(access_letter, FIRST_DATA_ROW, LAST_ROW, "Yes", "No"))

    last_col_letter = get_column_letter(len(HEADERS))
    table = Table(displayName="LoginInformationTable", ref=f"A{HEADER_ROW}:{last_col_letter}{FIRST_DATA_ROW}")
    table.tableStyleInfo = ds.modern_table_style()
    ws.add_table(table)

    # Quiet row banding + hairline row separators (§17.10), replacing the
    # built-in style's own banding/borders -- a fixed 50-row allowance, so
    # (unlike Archive/Campaigns Archive) direct per-cell borders are cheap
    # here.
    cf.apply_rules(ws, cf.row_banding_rule(f"A{FIRST_DATA_ROW}:{last_col_letter}{LAST_ROW}", FIRST_DATA_ROW))
    ds.apply_table_hairline_borders(ws, 1, len(HEADERS), FIRST_DATA_ROW, LAST_ROW)

    ws.sheet_properties.tabColor = ds.TAB_REFERENCE

    return ws
