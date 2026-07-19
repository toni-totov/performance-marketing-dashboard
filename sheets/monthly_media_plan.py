"""Monthly Media Plan -- mirrors the monthly campaign plan the user's
agency sends (PLAN.md §9): a color-coded campaign table, a per-week Gantt
timeline, a popup-driven monthly EUR->USD exchange rate, and a manual
month-reset flow. This module builds the *structure* only (Phase 2);
the exchange-rate popup and month-reset button are wired in Phase 3.

Layout, top to bottom:
  Rows 1-3: top stats strip -- 3 KPI cards (§17.10), each 2 columns wide:
    Planned Campaigns (A:B), Total Planned Budget EUR (D:E), Total
    Planned Budget USD (G:H). This is the first sheet to actually build
    the KPI-card convention described in §17.10 (Dashboard, which also
    uses it, is built later).
  Row 4: the exchange-rate button + live rate label, and the Start New
    Month button -- kept off to the side, away from the table (§9.4),
    both styled-only cells this phase (Phase 3 wires the click handlers).
  Row 5 (hidden): Gantt rolling anchor (PlanYear/PlanMonth-based, NOT
    TODAY()-based like Campaigns' rolling window -- this Gantt tracks the
    plan's own stored month, only advancing when ResetMonthlyPlan() runs)
    + the 6 raw week-start dates. 6 weeks always fully covers any month,
    same reasoning as Calendar's 6-row month grid (§8.2) -- but most
    months only need 4 or 5 of those 6 to actually be visible; see
    _needed_gantt_weeks() and modMonthlyMediaPlan.bas's
    UpdateGanttWeekVisibility for why the excess trailing column(s) are
    hidden rather than always shown (§19 Phase 5 correction).
  Row 6: visible header row -- Funded, Campaign Name, Platform,
    Optimization, Audience, Start Date, End Date, Budget (EUR),
    Budget (USD) (A-I), then 6 Gantt week-range labels (J-O).
  Row 7+: the Campaign Table (§9.2) -- a real Excel Table/ListObject
    (unlike Campaigns' reserved-slot plain-range layout), but **pre-sized
    to the full 30-row allowance (A6:I36) from the start**, per a
    live-tested finding that corrects two initial assumptions in a row:

    1. Writing a value into the cell directly below the table's last row
       does extend the ListObject boundary (confirmed via COM), but does
       NOT copy down that row's number format or formula from the row
       above -- the new row came back with `NumberFormat = "General"`
       and an empty Budget (USD) formula. Table-growth auto-fill-down is
       a UI-keystroke heuristic (Excel's "Extend Data Range Formats and
       Formulas" AutoCorrect option), not something a plain value write
       reliably triggers.
    2. The fix for #1 -- pre-writing the Budget (USD) formula and number
       formats into the not-yet-used rows below the table, same as the
       Gantt's own fixed allowance -- turned out to *break* native
       auto-expand-via-typing entirely: with real content already
       sitting in that row, Excel no longer treats it as the table's
       natural next-blank-row and refuses to extend the ListObject into
       it (confirmed via COM: growth succeeded when the row below was
       genuinely blank, failed once formulas/formats were pre-applied
       there).

    Given that conflict, pre-sizing the Table to all 30 rows up front
    sidesteps it completely: the Budget (USD) formula and every number
    format are simply live in every row from the start (no growth event
    to rely on at all), and this table's own dropdown/chip conditional
    formatting and data validation -- discrete objects with their own
    fixed sqref rather than a cell property, so never travel with Table
    growth regardless (same reasoning as archive.py) -- are pre-applied
    the same way. The Gantt (J:O) stays aligned 1:1 with the table's rows
    by construction, no VBA hook needed for any of it. If the user ever
    needs more than 30 campaigns in one month, Excel's native Insert Table
    Row still works past row 36; the Gantt/CF/DV simply wouldn't extend
    that far automatically.

Hidden column P holds PlanYear/PlanMonth/MonthlyExchangeRate, each a
workbook-scoped defined name. PlanYear/PlanMonth default to the current
year/month at build time (same pattern as Calendar's SelectedYear/
SelectedMonth) and are only updated by Phase 3's ResetMonthlyPlan() --
fully manual, no automatic month-change detection (§9.5).
MonthlyExchangeRate defaults to 1.08 (confirmed with the user) and is
carried over across month resets, only changed via frmExchangeRate.
"""

import datetime

from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table

from . import conditional_formatting as cf
from . import design_system as ds
from .week_grid import month_grid_anchor_formula, week_label_formula, week_number_formula, week_start_formula

SHEET_NAME = "Monthly Media Plan"

HIDDEN_DATE_ROW = 5
HEADER_ROW = 6
FIRST_DATA_ROW = 7
NUM_PREFORMATTED_ROWS = 30
GANTT_LAST_ROW = FIRST_DATA_ROW + NUM_PREFORMATTED_ROWS - 1  # 36
NUM_GANTT_WEEKS = 6

_DATE_FORMAT = "DD.MM.YYYY"
_EUR_FORMAT = '#,##0.00" €"'
_USD_FORMAT = '"$"#,##0.00'

HEADERS = [
    "Funded", "Campaign Name", "Platform", "Optimization", "Audience",
    "Start Date", "End Date", "Budget (EUR)", "Budget (USD)",
]
_COL = {name: i + 1 for i, name in enumerate(HEADERS)}
_GANTT_FIRST_COL = len(HEADERS) + 1  # J = 10
_GANTT_LAST_COL = _GANTT_FIRST_COL + NUM_GANTT_WEEKS - 1  # O = 15
_STATE_COL = get_column_letter(_GANTT_LAST_COL + 1)  # P = 16

_FUNDED_PAIRS = [("eStore", ds.CATEGORICAL[0]), ("MX", ds.CATEGORICAL[1]), ("Corporate", ds.CATEGORICAL[2])]
_PLATFORM_PAIRS = [("Google", ds.PLATFORM_GOOGLE), ("META", ds.PLATFORM_META)]
_OPTIMIZATION_PAIRS = [
    ("Performance Max", ds.CATEGORICAL[0]),
    ("Search", ds.CATEGORICAL[1]),
    ("Conversion (Purchases)", ds.CATEGORICAL[2]),
    ("Demand Gen", ds.CATEGORICAL[3]),
    ("Clicks", ds.CATEGORICAL[4]),
]

_DEFAULT_EXCHANGE_RATE = 1.08


def _needed_gantt_weeks(year, month):
    """How many of the 6 pre-built Gantt week columns actually overlap
    `year`/`month` (§19 Phase 5, user-requested/found live). The Gantt
    always builds 6 columns as a worst-case allowance (some month/
    weekday combinations genuinely need all 6 to fully cover the month,
    same reasoning as Calendar's 6-row grid) -- but most months only
    need 4 or 5, and the excess trailing column(s) would otherwise show
    days entirely in the FOLLOWING month, unrelated to the current plan
    (confirmed live: July 2026's 6th column was 100% August). Same
    Monday-start grid-anchor logic as week_grid.py's
    month_grid_anchor_formula, computed in Python since this needs to
    run at build time for the initial default -- modMonthlyMediaPlan.bas's
    UpdateGanttWeekVisibility is the VBA-side equivalent, needed because
    PlanYear/PlanMonth can also change at runtime via "Start New Month".
    """
    month_start = datetime.date(year, month, 1)
    if month == 12:
        month_end = datetime.date(year, 12, 31)
    else:
        month_end = datetime.date(year, month + 1, 1) - datetime.timedelta(days=1)
    grid_anchor = month_start - datetime.timedelta(days=month_start.weekday())  # Monday of that week

    needed = 0
    for w in range(NUM_GANTT_WEEKS):
        week_start = grid_anchor + datetime.timedelta(days=7 * w)
        week_end = week_start + datetime.timedelta(days=6)
        if week_end >= month_start and week_start <= month_end:
            needed = w + 1
    return needed


def build(wb):
    """Create the Monthly Media Plan sheet in `wb` (an openpyxl Workbook).
    Returns the created worksheet.
    """
    ws = wb.create_sheet(SHEET_NAME)
    ds.apply_modern_chrome(ws)
    ws.freeze_panes = ws.cell(row=FIRST_DATA_ROW, column=1)

    # --- Rows 1-3: KPI cards (§17.10) ---
    _build_kpi_card(
        ws, wb, caption_range="A1:B1", value_range="A2:B3",
        caption="Planned Campaigns",
        formula="=COUNTA(MonthlyMediaPlanTable[Campaign Name])",
        number_format="General", named_range=None,
    )
    _build_kpi_card(
        ws, wb, caption_range="D1:E1", value_range="D2:E3",
        caption="Total Planned Budget (EUR)",
        formula="=SUM(MonthlyMediaPlanTable[Budget (EUR)])",
        number_format=_EUR_FORMAT, named_range="TotalPlannedBudgetEUR",
    )
    _build_kpi_card(
        ws, wb, caption_range="G1:H1", value_range="G2:H3",
        caption="Total Planned Budget (USD)",
        formula="=SUM(MonthlyMediaPlanTable[Budget (USD)])",
        number_format=_USD_FORMAT, named_range="TotalPlannedBudgetUSD",
    )

    # --- Row 4: exchange-rate button + rate label, Start New Month button ---
    button_font = Font(name=ds.FONT_NAME, size=ds.FONT_SIZE_BODY, bold=True, color=ds.INK_PRIMARY)
    button_fill = PatternFill("solid", fgColor=ds.CARD_SURFACE)
    center = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A4:B4")
    rate_button = ws.cell(row=4, column=1, value=f"{ds.ICON_EXCHANGE_RATE} Set Exchange Rate")
    rate_button.font = button_font
    rate_button.fill = button_fill
    ds.apply_box_outline(ws, "A4:B4", color=ds.INK_MUTED)
    rate_button.alignment = center

    ws.merge_cells("C4:E4")
    # FIXED(), not TEXT(..., "0.00") -- live-tested finding: TEXT()'s
    # format-code argument is parsed against the current Excel locale's
    # own decimal separator, not a fixed period, and errors (#VALUE!) on
    # this Bulgarian-locale machine (decimal comma) when given a
    # period-decimal format code. FIXED() has no such format-code string
    # to mismatch and renders using whatever the current locale's decimal
    # separator actually is -- correct on any machine, not just this one.
    # 4 decimals, not 2 (user-requested, PLAN.md §19 Phase 5), matching
    # frmExchangeRate's own popup display.
    rate_label = ws.cell(row=4, column=3, value='="1 EUR = "&FIXED(MonthlyExchangeRate,4)&" USD"')
    rate_label.font = Font(name=ds.FONT_NAME, size=ds.FONT_SIZE_BODY, italic=True, color=ds.INK_SECONDARY)
    rate_label.alignment = center

    ws.merge_cells("G4:H4")
    reset_button = ws.cell(row=4, column=7, value=f"{ds.ICON_REFRESH} Start New Month")
    reset_button.font = button_font
    reset_button.fill = button_fill
    ds.apply_box_outline(ws, "G4:H4", color=ds.INK_MUTED)
    reset_button.alignment = center

    # --- Row 5 (hidden): Gantt anchor + raw week-start dates, driven by
    # PlanYear/PlanMonth (not TODAY()) -- see module docstring ---
    ws.cell(row=HIDDEN_DATE_ROW, column=1, value=month_grid_anchor_formula("PlanYear", "PlanMonth"))
    anchor_ref = f"A{HIDDEN_DATE_ROW}"
    for w in range(NUM_GANTT_WEEKS):
        col = _GANTT_FIRST_COL + w
        ws.cell(row=HIDDEN_DATE_ROW, column=col, value=week_start_formula(anchor_ref, w))
    ws.row_dimensions[HIDDEN_DATE_ROW].hidden = True

    # --- Row 6: visible header row ---
    for i, label in enumerate(HEADERS, start=1):
        cell = ws.cell(row=HEADER_ROW, column=i, value=label)
        ds.style_header_cell(cell)

    # "Week NN" + the date-range label, combined into one wrapped,
    # two-line cell -- same technique as Campaigns' Gantt (§19 Phase 5,
    # user-requested for consistency across both Gantts).
    ws.row_dimensions[HEADER_ROW].height = ds.HEADER_ROW_HEIGHT
    gap_border = Border(left=Side(style="medium", color=ds.INK_MUTED))
    for w in range(NUM_GANTT_WEEKS):
        col = _GANTT_FIRST_COL + w
        week_start_ref = f"{get_column_letter(col)}{HIDDEN_DATE_ROW}"
        week_end_ref = f"{week_start_ref}+6"
        week_number_expr = week_number_formula(week_start_ref)
        range_label_expr = week_label_formula(week_start_ref, week_end_ref)[1:]  # strip leading "="
        cell = ws.cell(row=HEADER_ROW, column=col, value=f'="Week "&{week_number_expr}&CHAR(10)&{range_label_expr}')
        ds.style_header_cell(cell)
        if w == 0:
            # Visual gap between the table and the Gantt (user-requested,
            # same fix as Campaigns) -- a medium left border on the
            # boundary column rather than an actual spacer column. Set
            # AFTER style_header_cell() so it overrides that cell's own
            # (thinner, muted) divider border with this stronger one.
            cell.border = gap_border

    # --- Column widths ---
    widths = {
        "A": 12, "B": 24, "C": 10, "D": 20, "E": 18, "F": 13, "G": 13, "H": 14, "I": 14,
    }
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width
    for w in range(NUM_GANTT_WEEKS):
        ws.column_dimensions[get_column_letter(_GANTT_FIRST_COL + w)].width = 9
    ws.column_dimensions[_STATE_COL].hidden = True

    # Hide any Gantt week column(s) that don't overlap the initial
    # PlanYear/PlanMonth default (today's month) at all -- see
    # _needed_gantt_weeks(). modMonthlyMediaPlan.bas's
    # UpdateGanttWeekVisibility keeps this correct after "Start New
    # Month" changes PlanYear/PlanMonth at runtime.
    _today = datetime.date.today()
    _needed = _needed_gantt_weeks(_today.year, _today.month)
    for w in range(NUM_GANTT_WEEKS):
        if w >= _needed:
            ws.column_dimensions[get_column_letter(_GANTT_FIRST_COL + w)].hidden = True

    # --- Number formats + the Budget (USD) formula, pre-applied across
    # the full fixed 30-row allowance (not just row 7) -- see module
    # docstring: Table auto-grow does NOT copy these down, confirmed via
    # live COM testing ---
    budget_eur_letter = get_column_letter(_COL["Budget (EUR)"])
    for row in range(FIRST_DATA_ROW, GANTT_LAST_ROW + 1):
        ws.cell(row=row, column=_COL["Start Date"]).number_format = _DATE_FORMAT
        ws.cell(row=row, column=_COL["End Date"]).number_format = _DATE_FORMAT
        ws.cell(row=row, column=_COL["Budget (EUR)"]).number_format = _EUR_FORMAT
        budget_usd_cell = ws.cell(row=row, column=_COL["Budget (USD)"])
        budget_usd_cell.value = (
            f'=IF({budget_eur_letter}{row}="","",{budget_eur_letter}{row}*MonthlyExchangeRate)'
        )
        budget_usd_cell.number_format = _USD_FORMAT
        # Visual gap between the table and the Gantt (user-requested,
        # same fix as Campaigns and the header row above).
        ws.cell(row=row, column=_GANTT_FIRST_COL).border = gap_border

    table = Table(
        displayName="MonthlyMediaPlanTable",
        ref=f"A{HEADER_ROW}:{get_column_letter(len(HEADERS))}{GANTT_LAST_ROW}",
    )
    table.tableStyleInfo = ds.modern_table_style()
    ws.add_table(table)

    # Quiet row banding + hairline row separators (§17.10) -- table
    # columns only (1..len(HEADERS)), not the Gantt columns, which keep
    # their own existing left-edge gap border untouched.
    cf.apply_rules(
        ws, cf.row_banding_rule(f"A{FIRST_DATA_ROW}:{get_column_letter(len(HEADERS))}{GANTT_LAST_ROW}", FIRST_DATA_ROW)
    )
    ds.apply_table_hairline_borders(ws, 1, len(HEADERS), FIRST_DATA_ROW, GANTT_LAST_ROW)

    # --- Dropdowns + categorical chip CF, pre-applied across the fixed
    # 30-row allowance (not just row 7) -- both are discrete objects with
    # their own fixed sqref, not a cell property Table-growth copies down
    # (same reasoning as Archive's CF ranges) ---
    _add_dropdown(ws, "Funded", '"eStore,MX,Corporate"')
    _add_dropdown(ws, "Platform", '"Google,META"')
    _add_dropdown(ws, "Optimization", '"Performance Max,Search,Conversion (Purchases),Demand Gen,Clicks"')

    cf.apply_rules(ws, cf.categorical_rules(get_column_letter(_COL["Funded"]), FIRST_DATA_ROW, GANTT_LAST_ROW, _FUNDED_PAIRS))
    cf.apply_rules(ws, cf.categorical_rules(get_column_letter(_COL["Platform"]), FIRST_DATA_ROW, GANTT_LAST_ROW, _PLATFORM_PAIRS))
    cf.apply_rules(ws, cf.categorical_rules(get_column_letter(_COL["Optimization"]), FIRST_DATA_ROW, GANTT_LAST_ROW, _OPTIMIZATION_PAIRS))

    # --- Gantt bars, colored by Platform (§9.3), one CF rule per week
    # column across the full fixed 30-row range -- a single contiguous
    # range per column, so (unlike Campaigns' 20 non-contiguous slots)
    # relative references retarget correctly per row without needing a
    # per-row loop ---
    start_letter = get_column_letter(_COL["Start Date"])
    end_letter = get_column_letter(_COL["End Date"])
    platform_letter = get_column_letter(_COL["Platform"])
    for w in range(NUM_GANTT_WEEKS):
        col = _GANTT_FIRST_COL + w
        col_letter = get_column_letter(col)
        week_start_ref = f"{col_letter}${HIDDEN_DATE_ROW}"
        week_end_ref = f"{week_start_ref}+6"
        gantt_range = f"{col_letter}{FIRST_DATA_ROW}:{col_letter}{GANTT_LAST_ROW}"
        start_cell = f"${start_letter}{FIRST_DATA_ROW}"
        end_cell = f"${end_letter}{FIRST_DATA_ROW}"
        platform_cell = f"${platform_letter}{FIRST_DATA_ROW}"
        for label, color in _PLATFORM_PAIRS:
            condition = f'AND({start_cell}<={week_end_ref},{end_cell}>={week_start_ref},{platform_cell}="{label}")'
            # bgColor, not fgColor -- see conditional_formatting.py's
            # _fill() docstring; this hand-rolled rule bypasses that
            # shared helper so needs the same fix applied directly.
            ws.conditional_formatting.add(gantt_range, FormulaRule(formula=[condition], fill=PatternFill("solid", bgColor=color)))

    # --- Hidden state: PlanYear/PlanMonth/MonthlyExchangeRate (column P) ---
    today = datetime.date.today()
    ws[f"{_STATE_COL}1"] = today.year
    ws[f"{_STATE_COL}2"] = today.month
    ws[f"{_STATE_COL}3"] = _DEFAULT_EXCHANGE_RATE

    _named(wb, "PlanYear", SHEET_NAME, f"{_STATE_COL}1")
    _named(wb, "PlanMonth", SHEET_NAME, f"{_STATE_COL}2")
    _named(wb, "MonthlyExchangeRate", SHEET_NAME, f"{_STATE_COL}3")

    # --- Sheet tab color (reference sheet, §17.9) ---
    ws.sheet_properties.tabColor = ds.TAB_REFERENCE

    return ws


def _build_kpi_card(ws, wb, caption_range, value_range, caption, formula, number_format, named_range):
    # §17.10 modernization pass: matches Dashboard's own KPI card redesign
    # (thick brand navy top accent, hairline sides/bottom, border
    # applied to every physical cell across both merges -- not just each
    # merge's anchor, which left the right-hand column unbordered in real
    # Excel; see design_system.apply_box_outline's docstring for why).
    ws.merge_cells(caption_range)
    caption_ref = caption_range.split(":")[0]
    caption_cell = ws[caption_ref]
    caption_cell.value = caption
    caption_cell.font = Font(name=ds.FONT_NAME, size=ds.FONT_SIZE_CAPTION, color=ds.INK_SECONDARY)
    caption_cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells(value_range)
    value_ref = value_range.split(":")[0]
    value_cell = ws[value_ref]
    value_cell.value = formula
    value_cell.number_format = number_format
    value_cell.font = Font(name=ds.FONT_NAME, size=ds.FONT_SIZE_KPI, bold=True, color=ds.INK_PRIMARY)
    value_cell.alignment = Alignment(horizontal="center", vertical="center")

    card_fill = PatternFill("solid", fgColor=ds.CARD_SURFACE)
    hair = Side(style="hair", color=ds.GRIDLINE)
    thick_top = Side(style="thick", color=ds.BRAND_ACCENT)
    cap_min_col, cap_min_row, cap_max_col, cap_max_row = range_boundaries(caption_range)
    val_min_col, val_min_row, val_max_col, val_max_row = range_boundaries(value_range)
    min_row, max_row = cap_min_row, val_max_row
    min_col, max_col = cap_min_col, cap_max_col
    for row in range(min_row, max_row + 1):
        for col in range(min_col, max_col + 1):
            cell = ws.cell(row=row, column=col)
            cell.fill = card_fill
            cell.border = Border(
                top=thick_top if row == min_row else None,
                bottom=hair if row == max_row else None,
                left=hair if col == min_col else None,
                right=hair if col == max_col else None,
            )

    if named_range:
        _named(wb, named_range, SHEET_NAME, value_ref)


def _add_dropdown(ws, column_name, formula1):
    dv = DataValidation(type="list", formula1=formula1, allow_blank=True)
    letter = get_column_letter(_COL[column_name])
    dv.add(f"{letter}{FIRST_DATA_ROW}:{letter}{GANTT_LAST_ROW}")
    ws.add_data_validation(dv)


def _named(wb, name, sheet_name, cell_ref):
    wb.defined_names[name] = DefinedName(name, attr_text=f"'{sheet_name}'!${cell_ref[0]}${cell_ref[1:]}")
