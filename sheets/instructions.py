"""Instructions -- a plain-language user guide (PLAN.md §16), written for
the end user rather than a technical reference. No VBA, no formulas
beyond the TOC's native internal hyperlinks. Structure:
  1. A prominent "Getting Started" callout (enabling macros) -- first
     thing on the sheet, since nothing else works without it.
  2. A table of contents with native Excel hyperlinks (no VBA) jumping to
     each of the 11 interactive sheets' own section below.
  3. One section per sheet: a title, a one-line "what it's for", and a
     handful of plain-language bullets describing its interactive
     behaviors. Content confirmed with the user via a sample section
     before writing all 11 (Tasks was the sample).

Row positions for every section are computed in a single pass up front
(each section's bullet count is fixed, known ahead of time), so the TOC's
jump targets can be written correctly without a two-pass build.
"""

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import design_system as ds

SHEET_NAME = "Instructions"

_GETTING_STARTED_TITLE = "Getting Started -- Enable Macros"
_GETTING_STARTED_BODY = (
    "When you first open this workbook, Excel shows a yellow security bar near the top saying "
    "something like \"SECURITY WARNING Macros have been disabled.\" Click \"Enable Content\" on "
    "that bar. This workbook won't work correctly until you do -- buttons, popups, and automatic "
    "updates all depend on macros being enabled."
)

# (title, "what it's for" line, [bullets])
_SECTIONS = [
    (
        "Dashboard",
        "Your home screen -- a live overview of tasks, campaigns, and productivity, pulled from every other sheet.",
        [
            'Click "Refresh Now" to force everything to recalculate immediately (it also refreshes automatically whenever you open the file).',
            "The KPI cards at the top update live as you work -- no action needed.",
            "Marketing Campaign Reminders shows any campaign starting or ending within the next 7 days.",
        ],
    ),
    (
        "Tasks",
        "Your main task tracker -- add tasks and subtasks, organize them into a hierarchy, and track deadlines, priority, and time spent.",
        [
            "Double-click a task's Deadline cell to pick a date from the Calendar.",
            'Double-click the Timer cell to start/stop tracking time on that task -- this also marks it "In Progress."',
            "Double-click a task's Notes cell to open the Notes popup and write a longer note.",
            "Double-click a task's name to collapse/expand its subtasks.",
            "Use the up/down arrow cells to reorder a task among its siblings.",
            'Set a task\'s Status to "Completed" to move it to Archive automatically.',
            'Set a task\'s Status to "Delete" to permanently remove it (and any subtasks) after a confirmation prompt -- this does not go to Archive.',
            "Pick a Linked Campaign from the dropdown to connect a task to a Campaign.",
        ],
    ),
    (
        "Archive",
        "A read-only historical record of every completed task -- filled in automatically, never edited by hand.",
        [
            "Double-click a task's Notes cell to view the note it had when it was completed.",
            'Tasks land here automatically when you set their Status to "Completed" on the Tasks sheet.',
            "This sheet is protected against direct edits, since Dashboard's historical stats read straight from it.",
        ],
    ),
    (
        "Calendar",
        "A date-picker popup, launched automatically whenever you double-click a Deadline cell elsewhere in the workbook.",
        [
            'Use the arrow buttons to move between months, or click "Today" to jump back to the current month.',
            "Double-click a day to select it, then set the time with the Hour/Minute dropdowns.",
            'Click "Confirm" to send the picked date back to the cell you started from, or "Cancel" to close without changes.',
        ],
    ),
    (
        "Campaigns",
        "Your own campaign tracker -- separate from the agency's Monthly Media Plan -- with linked tasks and a visual timeline.",
        [
            "Double-click a campaign's Notes cell to open the Notes popup.",
            "Double-click a campaign's name to collapse/expand its linked-tasks summary block (only works once it has at least one linked task).",
            "Use the up/down arrow cells to reorder a campaign.",
            "Every task you link to a campaign (via Tasks' Linked Campaign dropdown) shows up automatically in that campaign's summary block below it.",
            'Set a campaign\'s Status to "Completed" to move it to Campaigns Archive automatically.',
            "The colored bars show each campaign's timeline across the visible weeks.",
        ],
    ),
    (
        "Campaigns Archive",
        "A read-only historical record of every completed campaign -- filled in automatically, same idea as Archive.",
        [
            "Double-click a campaign's Notes cell to view its note.",
            'Campaigns land here automatically when marked "Completed" on the Campaigns sheet.',
        ],
    ),
    (
        "Monthly Media Plan",
        "The agency-provided monthly budget and schedule plan -- separate from your own Campaigns tracker.",
        [
            'Click "Set Exchange Rate" to update the EUR-to-USD rate used to calculate Budget (USD) -- it\'s shown just to the right of the button.',
            'Click "Start New Month" to clear the campaign table and move the Gantt to the current month -- this is destructive and asks for confirmation first.',
            "The Gantt bars are colored by Platform (Google/META).",
        ],
    ),
    (
        "Recurring Tasks",
        "Templates for tasks that repeat on a schedule (Daily/Weekly/Monthly) -- they generate real entries on Tasks automatically.",
        [
            'Set a template\'s Status to "Active" to start generating occurrences, or "Inactive" to pause it.',
            "Choose a Recurrence, then a Task Start On value (a weekday for Weekly, a day-of-month for Monthly).",
            "Missed occurrences catch up automatically the next time you open the workbook -- nothing is lost if you don't open it for a while.",
            "Double-click a template's Notes cell to open the Notes popup.",
        ],
    ),
    (
        "Schedule",
        "A visual monthly calendar showing your work location and meetings each day.",
        [
            'Use the arrow buttons to move between months, or click "Today" to jump back to the current month.',
            "Double-click a day to set its Work Location and manage meetings for that date.",
            "Each day's background color shows its Work Location; the small number below the date shows how many meetings are scheduled.",
            'Click "Reminder Settings" to change how many minutes before a meeting you\'re reminded, and how many reminders you get.',
        ],
    ),
    (
        "Login Information",
        "A flat reference list of platform logins -- type directly into most cells.",
        [
            "Passwords are stored in plain text here for simplicity -- this relies on your machine's own security, not Excel's.",
            'Set Access to "Yes" or "No" for each platform.',
            "Double-click a row's Notes cell to open the Notes popup.",
        ],
    ),
    (
        "Quick Links",
        "A flat reference list of bookmarked links -- type directly into the cells, no popups involved.",
        [
            "Excel automatically turns recognizable URLs into clickable links as you type them.",
        ],
    ),
]

_TITLE_ROW = 1
_GETTING_STARTED_TITLE_ROW = 3
_GETTING_STARTED_BODY_ROW = 4
_TOC_TITLE_ROW = 6
_TOC_FIRST_ROW = 7
_SECTION_GAP = 1  # blank row between sections


def build(wb):
    """Create the Instructions sheet in `wb` (an openpyxl Workbook).
    Returns the created worksheet.
    """
    ws = wb.create_sheet(SHEET_NAME)
    ds.apply_modern_chrome(ws)

    title_font = Font(name=ds.FONT_NAME, size=ds.FONT_SIZE_TITLE + 4, bold=True, color=ds.INK_PRIMARY)
    ws.cell(row=_TITLE_ROW, column=1, value="Instructions").font = title_font

    # --- Getting Started callout ---
    _write_callout(ws)

    # --- TOC heading ---
    toc_heading = ws.cell(row=_TOC_TITLE_ROW, column=1, value="Contents")
    toc_heading.font = Font(name=ds.FONT_NAME, size=ds.FONT_SIZE_TITLE, bold=True, color=ds.INK_PRIMARY)

    # --- Precompute every section's anchor row (single pass) ---
    section_rows = []
    row = _TOC_FIRST_ROW + len(_SECTIONS) + _SECTION_GAP + 1
    for title, blurb, bullets in _SECTIONS:
        section_rows.append(row)
        row += 2 + len(bullets) + _SECTION_GAP  # title + blurb + bullets + gap

    # --- TOC entries (native hyperlinks, no VBA) ---
    link_font = Font(name=ds.FONT_NAME, size=ds.FONT_SIZE_BODY, color=ds.BRAND_ACCENT, underline="single")
    for i, (title, _, _bullets) in enumerate(_SECTIONS):
        toc_row = _TOC_FIRST_ROW + i
        cell = ws.cell(row=toc_row, column=1, value=title)
        cell.hyperlink = f"#'{SHEET_NAME}'!A{section_rows[i]}"
        cell.font = link_font

    # --- Sections ---
    section_title_font = Font(name=ds.FONT_NAME, size=ds.FONT_SIZE_TITLE, bold=True, color=ds.BRAND_ACCENT_TEXT)
    section_title_fill = PatternFill("solid", fgColor=ds.BRAND_ACCENT)
    blurb_font = Font(name=ds.FONT_NAME, size=ds.FONT_SIZE_BODY, italic=True, color=ds.INK_SECONDARY)
    bullet_font = Font(name=ds.FONT_NAME, size=ds.FONT_SIZE_BODY, color=ds.INK_PRIMARY)

    for i, (title, blurb, bullets) in enumerate(_SECTIONS):
        anchor = section_rows[i]
        title_cell = ws.cell(row=anchor, column=1, value=title)
        title_cell.font = section_title_font
        title_cell.fill = section_title_fill
        ws.merge_cells(start_row=anchor, start_column=1, end_row=anchor, end_column=8)

        blurb_cell = ws.cell(row=anchor + 1, column=1, value=blurb)
        blurb_cell.font = blurb_font
        ws.merge_cells(start_row=anchor + 1, start_column=1, end_row=anchor + 1, end_column=8)

        for b_i, bullet in enumerate(bullets):
            bullet_row = anchor + 2 + b_i
            bullet_cell = ws.cell(row=bullet_row, column=1, value=f"• {bullet}")
            bullet_cell.font = bullet_font
            bullet_cell.alignment = Alignment(wrap_text=False)
            ws.merge_cells(start_row=bullet_row, start_column=1, end_row=bullet_row, end_column=8)

    ws.column_dimensions["A"].width = 110

    ws.sheet_properties.tabColor = ds.TAB_REFERENCE

    return ws


def _write_callout(ws):
    title_cell = ws.cell(row=_GETTING_STARTED_TITLE_ROW, column=1, value=_GETTING_STARTED_TITLE)
    title_cell.font = Font(name=ds.FONT_NAME, size=ds.FONT_SIZE_TITLE, bold=True, color=ds.BRAND_ACCENT_TEXT)
    title_cell.fill = PatternFill("solid", fgColor=ds.BRAND_ACCENT)
    ws.merge_cells(start_row=_GETTING_STARTED_TITLE_ROW, start_column=1, end_row=_GETTING_STARTED_TITLE_ROW, end_column=8)

    body_cell = ws.cell(row=_GETTING_STARTED_BODY_ROW, column=1, value=_GETTING_STARTED_BODY)
    body_cell.font = Font(name=ds.FONT_NAME, size=ds.FONT_SIZE_BODY, color=ds.INK_PRIMARY)
    body_cell.alignment = Alignment(wrap_text=True, vertical="top")
    body_cell.fill = PatternFill("solid", fgColor=ds.CARD_SURFACE)
    ws.merge_cells(start_row=_GETTING_STARTED_BODY_ROW, start_column=1, end_row=_GETTING_STARTED_BODY_ROW, end_column=8)
    ws.row_dimensions[_GETTING_STARTED_BODY_ROW].height = 45
