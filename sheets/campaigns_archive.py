"""Campaigns Archive -- where completed campaigns land automatically when
marked "Completed" on Campaigns (PLAN.md §11, §10.4). Flat historical log
only -- no Gantt, no linked-tasks rollup, no collapse logic (per explicit
"just an archive" instruction), consistent with how Tasks' own Archive is
also a flat log (§7). This module builds the *structure* only (Phase 2);
every value is a static snapshot written by the archiving macro at archive
time (Phase 3).

Column layout (§11.1), left to right:
  Campaign Name, Start Date, End Date, Status, Progress, Notes,
  Completed At
Hidden column, after Completed At:
  CampaignKey (retained so Tasks!Linked Campaign lookups keep resolving
  correctly after archiving)

`Status` is always "Completed" for every row here (no live formula, no
per-value conditional formatting needed -- it's a constant), so the
initial data row's Status cell just gets a static workflow-status "done"
fill (§17.4) directly rather than a conditional-formatting rule, and VBA's
ListRows.Add (Phase 3) inherits that formatting as the table grows, same
as every other formatted column here.

`Campaigns Archive` is a genuine auto-expanding Excel Table, like Time Log
and Tasks' Archive -- rows are append-only and never reordered (§11.2), so
it starts at the minimum valid size and grows via VBA's ListRows.Add.

The whole sheet is protected (locked against direct cell edits), consistent
with the same decision already made for Tasks' Archive (§7.1) -- Dashboard's
Completed Campaigns Count (§11.3) reads straight from this data. Only basic
protection is set here; `UserInterfaceOnly:=True` does not persist across
file close/reopen and must be re-asserted by Phase 3's Workbook_Open, same
as Archive.
"""

from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table

from . import conditional_formatting as cf
from . import design_system as ds

SHEET_NAME = "Campaigns Archive"
HEADER_ROW = 1
FIRST_DATA_ROW = 2

_DATE_FORMAT = "DD.MM.YYYY"
_DATE_TIME_FORMAT = "DD.MM.YYYY HH:MM"
_PERCENT_FORMAT = "0%"

HEADERS = [
    "Campaign Name",
    "Start Date",
    "End Date",
    "Status",
    "Progress",
    "Notes",
    "Completed At",
]
HIDDEN_HEADERS = ["CampaignKey"]

_COL = {name: i + 1 for i, name in enumerate(HEADERS)}
_HIDDEN_COL = {name: len(HEADERS) + i + 1 for i, name in enumerate(HIDDEN_HEADERS)}


def build(wb):
    """Create the Campaigns Archive sheet in `wb` (an openpyxl Workbook).
    Returns the created worksheet.
    """
    ws = wb.create_sheet(SHEET_NAME)
    ds.apply_modern_chrome(ws)
    ws.freeze_panes = ws.cell(row=FIRST_DATA_ROW, column=1)
    ws.row_dimensions[HEADER_ROW].height = ds.HEADER_ROW_HEIGHT

    for col, label in enumerate([*HEADERS, *HIDDEN_HEADERS], start=1):
        cell = ws.cell(row=HEADER_ROW, column=col, value=label)
        ds.style_header_cell(cell)

    ws.column_dimensions[get_column_letter(_HIDDEN_COL["CampaignKey"])].hidden = True

    # Completed At: 19, not 15 (§17.10 bugfix, same zoom-vs-column-width
    # regression as Archive/Tasks). Start/End Date stay at 13 -- plain
    # "DD.MM.YYYY" (10 chars) still has enough margin at 115% zoom, unlike
    # the 16-character "DD.MM.YYYY HH:MM" format.
    widths = {
        "Campaign Name": 22,
        "Start Date": 13,
        "End Date": 13,
        "Status": 12,
        "Progress": 10,
        "Notes": 8,
        "Completed At": 19,
    }
    for name, width in widths.items():
        ws.column_dimensions[get_column_letter(_COL[name])].width = width

    # Number formats on the initial data row (row 2) -- VBA's ListRows.Add
    # (Phase 3) inherits formatting from the row above as the table grows,
    # same pattern already verified working for Time Log and Archive.
    ws.cell(row=FIRST_DATA_ROW, column=_COL["Start Date"]).number_format = _DATE_FORMAT
    ws.cell(row=FIRST_DATA_ROW, column=_COL["End Date"]).number_format = _DATE_FORMAT
    ws.cell(row=FIRST_DATA_ROW, column=_COL["Progress"]).number_format = _PERCENT_FORMAT
    ws.cell(row=FIRST_DATA_ROW, column=_COL["Notes"]).number_format = ds.notes_icon_format()
    ws.cell(row=FIRST_DATA_ROW, column=_COL["Completed At"]).number_format = _DATE_TIME_FORMAT

    # Status is a constant ("Completed" for every row here) -- a static
    # workflow-status "done" fill (§17.4) rather than a conditional-
    # formatting rule, since there's no varying value to key off.
    ws.cell(row=FIRST_DATA_ROW, column=_COL["Status"]).fill = PatternFill("solid", fgColor=ds.WORKFLOW_DONE)

    last_col_letter = get_column_letter(len(HEADERS) + len(HIDDEN_HEADERS))
    table = Table(displayName="CampaignsArchiveTable", ref=f"A1:{last_col_letter}{FIRST_DATA_ROW}")
    table.tableStyleInfo = ds.modern_table_style()
    ws.add_table(table)

    # Quiet row banding (§17.10), same generous growth-range reasoning as
    # Archive -- CF-based, no per-cell border bloat on an unbounded table.
    _CF_LAST_ROW = 1000
    cf.apply_rules(ws, cf.row_banding_rule(f"A{FIRST_DATA_ROW}:{last_col_letter}{_CF_LAST_ROW}", FIRST_DATA_ROW))

    # --- Sheet protection: locked against direct edits, consistent with
    # Tasks' Archive (§7.1 correction). UserInterfaceOnly must be
    # re-asserted by Phase 3's Workbook_Open -- see module docstring.
    ws.protection.sheet = True

    # --- Sheet tab color (archive sheet, §17.9) ---
    ws.sheet_properties.tabColor = ds.TAB_ARCHIVE

    return ws
