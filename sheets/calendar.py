"""Calendar -- the destination of the Deadline double-click handoff from
Tasks (PLAN.md §8). This module builds the *structure* only (Phase 2):
the month grid, navigation/button cells, Hour/Minute pickers, the live
preview label, and the handoff cells. Click-wiring (Prev/Next/Today,
double-click date selection, Confirm/Cancel) is VBA, added in Phase 3.

Layout (7 columns, A-G, matching the day grid):
  Row 1  (banner): <Prev | month/year label (merged B1:E1) | Today | Next>
  Row 2: weekday headers, Mon-Sun
  Rows 3-8: the 6x7 day grid (week_grid.py's formulas)
  Row 10: Hour label+dropdown, Minute label+dropdown
  Row 11: live preview label (merged A11:G11)
  Row 12: Confirm button (merged A12:C12) | Cancel button (merged E12:G12)
  Column J (hidden): SelectedYear/SelectedMonth/SelectedDate/
    SelectedDateTime as named cells, same pattern as the handoff cells
  Column Z (hidden): the 4 generic Calendar-handoff cells (§8.1),
    built by calendar_handoff.write_handoff_cells()

Implementation refinement over the original §8.3 wording: "selected date"
highlighting is a conditional-formatting rule (keyed off the named
SelectedDate cell), not VBA manually painting/clearing a fill on each
click. This is strictly simpler for Phase 3 -- the double-click handler
only needs to set SelectedDate's value; the highlight (and its automatic
clearing from whichever cell was previously selected) follows for free.
"""

import datetime

from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from . import design_system as ds
from .calendar_handoff import write_handoff_cells
from .week_grid import (
    month_grid_anchor_formula,
    month_grid_formulas,
    month_name_full_formula,
    weekday_name_formula,
)

SHEET_NAME = "Calendar"

_WEEKDAY_HEADERS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
_GRID_FIRST_ROW = 3
_GRID_LAST_ROW = 8  # 6 weeks
_WEEKEND_COLS = ("F", "G")  # Sat, Sun under a Monday-start layout

_STATE_COL = "J"  # hidden utility column for Calendar's own named state cells


def _named(wb, name, sheet_name, cell_ref):
    wb.defined_names[name] = DefinedName(name, attr_text=f"'{sheet_name}'!${cell_ref[0]}${cell_ref[1:]}")


def build(wb):
    """Create the Calendar sheet in `wb` (an openpyxl Workbook). Returns
    the created worksheet.
    """
    ws = wb.create_sheet(SHEET_NAME)
    ds.apply_modern_chrome(ws)

    brand_fill = PatternFill("solid", fgColor=ds.BRAND_ACCENT)
    white_bold = Font(name=ds.FONT_NAME, color=ds.BRAND_ACCENT_TEXT, bold=True)
    center = Alignment(horizontal="center", vertical="center")

    # --- Row 1: banner (Prev / month-year label / Today / Next) ---
    ws.row_dimensions[1].height = 26
    for col in range(1, 8):
        ws.cell(row=1, column=col).fill = brand_fill

    prev_cell = ws.cell(row=1, column=1, value=ds.ICON_NAV_PREV)
    prev_cell.font = Font(name=ds.FONT_NAME, color=ds.BRAND_ACCENT_TEXT, bold=True, size=ds.FONT_SIZE_TITLE)
    prev_cell.alignment = center

    ws.merge_cells("B1:E1")
    label_cell = ws.cell(row=1, column=2)
    # Locale-proof month name (PLAN.md §19 Phase 5, found live) -- see
    # week_grid.py's module-level comment for why TEXT(...,"MMMM") isn't
    # safe (renders in the workbook's system locale).
    _month_anchor = "DATE(SelectedYear,SelectedMonth,1)"
    label_cell.value = f'={month_name_full_formula(_month_anchor)}&" "&TEXT({_month_anchor},"YYYY")'
    label_cell.font = Font(name=ds.FONT_NAME, color=ds.BRAND_ACCENT_TEXT, bold=True, size=ds.FONT_SIZE_TITLE)
    label_cell.alignment = center

    today_cell = ws.cell(row=1, column=6, value="Today")
    today_cell.font = white_bold
    today_cell.alignment = center

    next_cell = ws.cell(row=1, column=7, value=ds.ICON_NAV_NEXT)
    next_cell.font = Font(name=ds.FONT_NAME, color=ds.BRAND_ACCENT_TEXT, bold=True, size=ds.FONT_SIZE_TITLE)
    next_cell.alignment = center

    # --- Row 2: weekday headers -- divider between columns (§17.11) ---
    for i, label in enumerate(_WEEKDAY_HEADERS):
        cell = ws.cell(row=2, column=i + 1, value=label)
        ds.style_header_cell(cell, wrap=False, divider=(i < len(_WEEKDAY_HEADERS) - 1))

    # --- Rows 3-8: day grid ---
    anchor_cell_ref = f"{_STATE_COL}6"  # a scratch cell for the grid anchor formula
    ws[anchor_cell_ref] = month_grid_anchor_formula("SelectedYear", "SelectedMonth")
    grid_formulas = month_grid_formulas(anchor_cell_ref)

    day_font = Font(name=ds.FONT_NAME, size=ds.FONT_SIZE_BODY)
    weekend_fill = PatternFill("solid", fgColor=ds.GRIDLINE)  # light neutral tint

    for r_offset, row_formulas in enumerate(grid_formulas):
        row = _GRID_FIRST_ROW + r_offset
        ws.row_dimensions[row].height = 22
        for c_offset, formula in enumerate(row_formulas):
            col = c_offset + 1
            cell = ws.cell(row=row, column=col, value=formula)
            cell.number_format = "D"
            cell.font = day_font
            cell.alignment = center
            letter = get_column_letter(col)
            if letter in _WEEKEND_COLS:
                cell.fill = weekend_fill

    grid_range = f"A{_GRID_FIRST_ROW}:G{_GRID_LAST_ROW}"
    top_left = f"A{_GRID_FIRST_ROW}"

    # Adjacent-month days: muted font (does not conflict with fill/border rules below)
    ws.conditional_formatting.add(
        grid_range,
        FormulaRule(formula=[f"MONTH({top_left})<>SelectedMonth"], font=Font(color=ds.INK_MUTED)),
    )
    # Days with a matching Tasks!Deadline: light accent fill
    # bgColor, not fgColor -- see conditional_formatting.py's _fill()
    # docstring; this hand-rolled rule bypasses that shared helper so
    # needs the same fix applied directly.
    ws.conditional_formatting.add(
        grid_range,
        FormulaRule(
            formula=[f'COUNTIFS(Tasks!$K:$K,">="&{top_left},Tasks!$K:$K,"<"&{top_left}+1)>0'],
            fill=PatternFill("solid", bgColor=ds.SEQUENTIAL_BLUE[150]),
        ),
    )
    # Selected date: distinct border, added BEFORE "today" so it takes
    # precedence if a cell is both today's date and the current selection.
    selected_border = Border(*[Side(style="medium", color=ds.CATEGORICAL[4])] * 4)  # violet
    ws.conditional_formatting.add(
        grid_range, FormulaRule(formula=[f"{top_left}=SelectedDate"], border=selected_border)
    )
    # Today: brand navy border
    today_border = Border(*[Side(style="medium", color=ds.BRAND_ACCENT)] * 4)
    ws.conditional_formatting.add(grid_range, FormulaRule(formula=[f"{top_left}=TODAY()"], border=today_border))

    # --- Rows 10-11: Hour/Minute pickers + live preview, grouped into one
    # card (§17.10 correction, user-reported "feels off, in the middle of
    # nowhere" -- these two rows had no background/border of their own
    # at all, floating between the day grid above and the Confirm/Cancel
    # buttons below with nothing visually tying them together). ---
    card_fill = PatternFill("solid", fgColor=ds.CARD_SURFACE)
    for row in (10, 11):
        for col in range(1, 8):
            ws.cell(row=row, column=col).fill = card_fill
    ds.apply_box_outline(ws, "A10:G11")

    ws.cell(row=10, column=1, value="Hour:").font = Font(name=ds.FONT_NAME, bold=True)
    hour_cell = ws.cell(row=10, column=2, value="00")
    ws.cell(row=10, column=3, value="Minute:").font = Font(name=ds.FONT_NAME, bold=True)
    minute_cell = ws.cell(row=10, column=4, value="00")

    hour_list = ",".join(f"{h:02d}" for h in range(24))
    hour_dv = DataValidation(type="list", formula1=f'"{hour_list}"', allow_blank=False)
    hour_dv.add(hour_cell.coordinate)
    ws.add_data_validation(hour_dv)

    minute_dv = DataValidation(type="list", formula1='"00,15,30,45"', allow_blank=False)
    minute_dv.add(minute_cell.coordinate)
    ws.add_data_validation(minute_dv)

    # --- Row 11: live preview label ---
    ws.merge_cells("A11:G11")
    preview_cell = ws.cell(row=11, column=1)
    # Locale-proof weekday/month names (PLAN.md §19 Phase 5, found live) --
    # see week_grid.py's module-level comment for why TEXT() with
    # "dddd"/"MMMM" isn't safe.
    _preview_date_expr = (
        f'{weekday_name_formula("SelectedDateTime")}&", "&TEXT(SelectedDateTime,"DD")'
        f'&" "&{month_name_full_formula("SelectedDateTime")}&" "&TEXT(SelectedDateTime,"YYYY")'
    )
    preview_cell.value = (
        f'=IF(SelectedDate="","No date selected yet",'
        f'{_preview_date_expr}&" — "&TEXT(SelectedDateTime,"HH:MM"))'
    )
    preview_cell.font = Font(name=ds.FONT_NAME, italic=True, color=ds.INK_SECONDARY)
    preview_cell.alignment = center

    # --- Row 12: Confirm / Cancel buttons ---
    ws.merge_cells("A12:C12")
    confirm_cell = ws.cell(row=12, column=1, value=f"{ds.ICON_CONFIRM} Confirm")
    confirm_cell.fill = brand_fill
    confirm_cell.font = white_bold
    confirm_cell.alignment = center

    ws.merge_cells("E12:G12")
    cancel_cell = ws.cell(row=12, column=5, value=f"{ds.ICON_CANCEL} Cancel")
    cancel_cell.font = Font(name=ds.FONT_NAME, bold=True, color=ds.INK_PRIMARY)
    cancel_cell.fill = PatternFill("solid", fgColor=ds.CARD_SURFACE)
    ds.apply_box_outline(ws, "E12:G12", color=ds.INK_MUTED)
    cancel_cell.alignment = center

    # --- Column widths ---
    for col in range(1, 8):
        ws.column_dimensions[get_column_letter(col)].width = 10

    # --- Hidden state cells (column J) + named ranges ---
    today = datetime.date.today()
    ws[f"{_STATE_COL}1"] = today.year
    ws[f"{_STATE_COL}2"] = today.month
    ws[f"{_STATE_COL}3"] = ""  # SelectedDate -- blank until a double-click sets it (Phase 3)
    ws[f"{_STATE_COL}4"] = (
        f'=IF({_STATE_COL}3="","",{_STATE_COL}3+TIME(VALUE(SelectedHour),VALUE(SelectedMinute),0))'
    )
    ws.column_dimensions[_STATE_COL].hidden = True

    _named(wb, "SelectedYear", SHEET_NAME, f"{_STATE_COL}1")
    _named(wb, "SelectedMonth", SHEET_NAME, f"{_STATE_COL}2")
    _named(wb, "SelectedDate", SHEET_NAME, f"{_STATE_COL}3")
    _named(wb, "SelectedDateTime", SHEET_NAME, f"{_STATE_COL}4")
    _named(wb, "SelectedHour", SHEET_NAME, "B10")
    _named(wb, "SelectedMinute", SHEET_NAME, "D10")

    # --- Generic Calendar handoff cells (§8.1, Phase 1) ---
    write_handoff_cells(wb, SHEET_NAME)

    # --- Sheet tab color (core daily-use sheet, §17.9) ---
    ws.sheet_properties.tabColor = ds.TAB_CORE

    return ws
