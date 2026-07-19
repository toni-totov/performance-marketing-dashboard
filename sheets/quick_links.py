"""Quick Links -- a simple flat reference table (PLAN.md §15), same
lightweight treatment as Login Information (§14): no VBA, no hierarchy,
no formulas, plain free-text Notes (not the frmNotes popup). A genuine
auto-expanding Excel Table; Excel auto-hyperlinks recognizable URLs typed
into the Link/URL column natively, no code needed.

Row allowance: 50 rows, confirmed with the user (same reasoning as Login
Information -- CF/DV would be pre-applied over a fixed range too, but
this sheet has no dropdown/chip field at all, so it's moot here).
"""

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table

from . import conditional_formatting as cf
from . import design_system as ds

SHEET_NAME = "Quick Links"
HEADER_ROW = 1
FIRST_DATA_ROW = 2
LAST_ROW = FIRST_DATA_ROW + 50 - 1  # 51

HEADERS = ["Category / Group", "Link / URL", "Description", "Notes"]
_COL = {name: i + 1 for i, name in enumerate(HEADERS)}


def build(wb):
    """Create the Quick Links sheet in `wb` (an openpyxl Workbook).
    Returns the created worksheet.
    """
    ws = wb.create_sheet(SHEET_NAME)
    ds.apply_modern_chrome(ws)
    ws.freeze_panes = ws.cell(row=FIRST_DATA_ROW, column=1)
    ws.row_dimensions[HEADER_ROW].height = ds.HEADER_ROW_HEIGHT

    for col, label in enumerate(HEADERS, start=1):
        cell = ws.cell(row=HEADER_ROW, column=col, value=label)
        ds.style_header_cell(cell)

    widths = {
        "Category / Group": 18,
        "Link / URL": 32,
        "Description": 28,
        "Notes": 24,
    }
    for name, width in widths.items():
        ws.column_dimensions[get_column_letter(_COL[name])].width = width

    last_col_letter = get_column_letter(len(HEADERS))
    table = Table(displayName="QuickLinksTable", ref=f"A{HEADER_ROW}:{last_col_letter}{FIRST_DATA_ROW}")
    table.tableStyleInfo = ds.modern_table_style()
    ws.add_table(table)

    # Quiet row banding + hairline row separators (§17.10).
    cf.apply_rules(ws, cf.row_banding_rule(f"A{FIRST_DATA_ROW}:{last_col_letter}{LAST_ROW}", FIRST_DATA_ROW))
    ds.apply_table_hairline_borders(ws, 1, len(HEADERS), FIRST_DATA_ROW, LAST_ROW)

    ws.sheet_properties.tabColor = ds.TAB_REFERENCE

    return ws
