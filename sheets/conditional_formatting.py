"""Reusable conditional-formatting rule generators for the color schemes
defined in PLAN.md §17: Priority/urgency severity coloring (§17.3) and the
workflow-status mini-scheme (§17.4). Applied first on Tasks; Recurring
Tasks and other sheets with the same fields reuse these rather than each
sheet re-deriving the same rule logic.

Each function returns a list of (range_string, FormulaRule) tuples --
apply them to a worksheet via apply_rules(ws, rules). The color fill sits
behind already-visible dropdown text (e.g. "Low"/"Medium"/"High" is the
cell's own text), so this satisfies §17.3's "never color alone" rule
inherently -- the color is a supporting cue, not the only signal.
"""

from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill

from . import design_system as ds


def _fill(hex_color):
    # bgColor, not fgColor -- live-tested finding (Monthly Media Plan
    # phase): for a *differential* (conditional-formatting) PatternFill,
    # Excel reads the solid fill's displayed color from bgColor. fgColor
    # is silently ignored (Interior.Color reads back as 0/black),
    # confirmed via COM by comparing fgColor-only, bgColor-only, and
    # both-set variants of the identical fill. This is the opposite of a
    # normal cell style's PatternFill, where fgColor is the one that
    # matters -- a genuine openpyxl/Excel dxf quirk, not a typo.
    return PatternFill("solid", bgColor=hex_color)


def _chip_font(hex_color):
    """White bold text on a dark chip, default (dark) text on a light one
    (V4, PLAN.md §22.8 -- dark text on the grey "Not Started", blue "In
    Progress", red "High" and green "Active" chips was hard to read).
    Decided by the fill's relative luminance, so any future chip color
    gets the right text automatically. 0.3 rather than the strict WCAG
    crossover (~0.18): on these mid-tone chips white reads clearly better
    in practice even where black's ratio is marginally higher.
    """
    def channel(c):
        c = c / 255
        return c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4

    r, g, b = (int(hex_color[i:i + 2], 16) for i in (0, 2, 4))
    lum = 0.2126 * channel(r) + 0.7152 * channel(g) + 0.0722 * channel(b)
    return Font(color=ds.BRAND_ACCENT_TEXT, bold=True) if lum < 0.3 else None


def _chip(formula, hex_color):
    return FormulaRule(formula=[formula], fill=_fill(hex_color), font=_chip_font(hex_color))


def apply_rules(ws, rules):
    """Add every (range_string, rule) pair in `rules` to `ws`."""
    for range_string, rule in rules:
        ws.conditional_formatting.add(range_string, rule)


def priority_rules(column_letter, first_row, last_row):
    """§17.3: Priority severity coloring -- Low=good, Medium=warning,
    High=critical.
    """
    rng = f"{column_letter}{first_row}:{column_letter}{last_row}"
    cell = f"${column_letter}{first_row}"
    return [
        (rng, _chip(f'{cell}="Low"', ds.PRIORITY_LOW)),
        (rng, _chip(f'{cell}="Medium"', ds.PRIORITY_MEDIUM)),
        (rng, _chip(f'{cell}="High"', ds.PRIORITY_HIGH)),
    ]


def urgency_rules(column_letter, first_row, last_row):
    """§17.3: urgency/deadline-proximity coloring, for a numeric "days
    remaining"-style column (e.g. Days Left): overdue (<0)=critical,
    due-soon (0-2 inclusive)=warning, on-track (>2)=good.
    """
    rng = f"{column_letter}{first_row}:{column_letter}{last_row}"
    cell = f"${column_letter}{first_row}"
    # ISNUMBER guard (PLAN.md §19 Phase 5, found live): a formula-based CF
    # numeric comparison coerces a blank/"" cell (e.g. an IFERROR(...,"")
    # fallback on an unused summary-row slot) to 0, so ">=0,<=2" was
    # silently matching and coloring cells that should stay uncolored.
    return [
        (rng, FormulaRule(formula=[f"AND(ISNUMBER({cell}),{cell}<0)"], fill=_fill(ds.URGENCY_OVERDUE))),
        (rng, FormulaRule(formula=[f"AND(ISNUMBER({cell}),{cell}>=0,{cell}<=2)"], fill=_fill(ds.URGENCY_DUE_SOON))),
        (rng, FormulaRule(formula=[f"AND(ISNUMBER({cell}),{cell}>2)"], fill=_fill(ds.URGENCY_ON_TRACK))),
    ]


def deadline_urgency_rules(column_letter, first_row, last_row):
    """§17.3: the same overdue/due-soon/on-track severity scheme as
    urgency_rules(), but driven directly by a DATE column compared
    against TODAY() rather than a precomputed "days remaining" number --
    for a column that shows the deadline itself (e.g. Campaigns' linked-
    task summary rows, PLAN.md §19 Phase 5) with no separate days-left
    value stored alongside it. Same ISNUMBER guard and same reasoning
    (blank/"" cells must not be coerced into a match).
    """
    rng = f"{column_letter}{first_row}:{column_letter}{last_row}"
    cell = f"${column_letter}{first_row}"
    return [
        (rng, FormulaRule(formula=[f"AND(ISNUMBER({cell}),{cell}<TODAY())"], fill=_fill(ds.URGENCY_OVERDUE))),
        (rng, FormulaRule(formula=[f"AND(ISNUMBER({cell}),{cell}>=TODAY(),{cell}<=TODAY()+2)"], fill=_fill(ds.URGENCY_DUE_SOON))),
        (rng, FormulaRule(formula=[f"AND(ISNUMBER({cell}),{cell}>TODAY()+2)"], fill=_fill(ds.URGENCY_ON_TRACK))),
    ]


def workflow_status_rules(
    column_letter,
    first_row,
    last_row,
    not_started_label="Not Started",
    in_motion_label="In Progress",
    done_label="Completed",
):
    """§17.4: the 3-state workflow-status scheme, reused across every
    lifecycle Status field (Tasks' Not Started/In Progress, Campaigns'
    Planned/Active/Completed, etc.) so "what stage is this in" reads
    identically everywhere. Pass whichever label text a given sheet
    actually uses for each stage -- the color mapping stays fixed.
    """
    rng = f"{column_letter}{first_row}:{column_letter}{last_row}"
    cell = f"${column_letter}{first_row}"
    return [
        (rng, _chip(f'{cell}="{not_started_label}"', ds.WORKFLOW_NOT_STARTED)),
        (rng, _chip(f'{cell}="{in_motion_label}"', ds.WORKFLOW_IN_MOTION)),
        (rng, _chip(f'{cell}="{done_label}"', ds.WORKFLOW_DONE)),
    ]


def diverging_rules(column_letter, first_row, last_row, dead_zone=None):
    """§17.5: diverging color scale for a genuine polarity metric (e.g.
    Archive's frozen `Days Left`, "days early/late at completion"):
    positive (early) = blue, negative (late) = red, exactly zero
    (on time) = neutral gray midpoint.

    `dead_zone` (V4 Archive "Early / Late", PLAN.md §22.5): values within
    +/- dead_zone count as the midpoint, matching a number format that
    prints them as "on time", and blank cells (e.g. a Skipped row) get no
    color at all instead of reading as 0.
    """
    rng = f"{column_letter}{first_row}:{column_letter}{last_row}"
    cell = f"${column_letter}{first_row}"
    if dead_zone is not None:
        num = f"ISNUMBER({cell})"
        return [
            (rng, FormulaRule(formula=[f"AND({num},{cell}>={dead_zone})"], fill=_fill(ds.DIVERGING_POSITIVE))),
            (rng, FormulaRule(formula=[f"AND({num},{cell}<=-{dead_zone})"], fill=_fill(ds.DIVERGING_NEGATIVE))),
            (rng, FormulaRule(formula=[f"AND({num},ABS({cell})<{dead_zone})"], fill=_fill(ds.DIVERGING_MIDPOINT))),
        ]
    return [
        (rng, FormulaRule(formula=[f"{cell}>0"], fill=_fill(ds.DIVERGING_POSITIVE))),
        (rng, FormulaRule(formula=[f"{cell}<0"], fill=_fill(ds.DIVERGING_NEGATIVE))),
        (rng, FormulaRule(formula=[f"{cell}=0"], fill=_fill(ds.DIVERGING_MIDPOINT))),
    ]


def categorical_rules(column_letter, first_row, last_row, label_color_pairs):
    """§17.2: identity-chip coloring for a dropdown field with its own
    fixed categorical slot assignment (e.g. Monthly Media Plan's Funded/
    Platform/Optimization). `label_color_pairs` is an ordered list of
    (label, hex_color) tuples -- pass the field's own values paired with
    their assigned `design_system.CATEGORICAL` slots, in the order
    §17.2 assigns them. Shared so every dropdown field with this
    treatment (this sheet now, Recurring Tasks'/Schedule's Recurrence
    and Location later) derives its rules the same way.
    """
    rng = f"{column_letter}{first_row}:{column_letter}{last_row}"
    cell = f"${column_letter}{first_row}"
    return [
        (rng, _chip(f'{cell}="{label}"', color))
        for label, color in label_color_pairs
    ]


def sequential_rules(column_letter, first_row, last_row, buckets):
    """§17.5: bucket a numeric column into the sequential blue ramp,
    light->dark = low->high magnitude/proportion (Campaigns `Progress`,
    Tasks/Recurring Tasks `Number of Subtasks` -- PLAN.md §19 Phase 6).

    `buckets` is an ascending list of (lower_bound_inclusive,
    sequential_step) pairs -- a cell matches the *last* bucket whose
    lower_bound it's >= to, open-ended at the top (so the same list shape
    works for a bounded 0..1 fraction like Progress % and an open-ended
    count like Number of Subtasks). Same ISNUMBER guard as urgency_rules
    -- a value below the first bucket's lower_bound, or a genuinely
    non-numeric "" (an IFERROR/IF(...,"") fallback on an unused row),
    gets no fill at all.
    """
    rng = f"{column_letter}{first_row}:{column_letter}{last_row}"
    cell = f"${column_letter}{first_row}"
    rules = []
    for i, (lower_bound, step) in enumerate(buckets):
        if i + 1 < len(buckets):
            next_lower = buckets[i + 1][0]
            formula = f"AND(ISNUMBER({cell}),{cell}>={lower_bound},{cell}<{next_lower})"
        else:
            formula = f"AND(ISNUMBER({cell}),{cell}>={lower_bound})"
        rules.append((rng, FormulaRule(formula=[formula], fill=_fill(ds.SEQUENTIAL_BLUE[step]))))
    return rules


def row_banding_rule(range_string, first_row, stripe_color=None):
    """§17.10 (PLAN.md §19 Phase 6, user-requested modernization pass): a
    single quiet alternating-row fill for a table body, replacing the
    built-in `TableStyleMedium2` banding every ListObject Table used to
    carry (see `design_system.modern_table_style()`'s docstring for why
    that was removed). `first_row` anchors the `MOD(ROW(),2)` parity check
    to the range's own first data row, not an absolute row 1, so a table
    starting on any row always begins its banding on the same visual
    parity regardless of where it sits on the sheet.
    """
    color = stripe_color or ds.TABLE_STRIPE
    return [(range_string, FormulaRule(formula=[f"MOD(ROW()-{first_row},2)=0"], fill=_fill(color)))]


def is_banding_rule(rule):
    """True for a rule produced by row_banding_rule() (MOD(ROW()-n,2)=0)."""
    return any(f.replace(" ", "").startswith("MOD(ROW()-") for f in (rule.formula or []))


def demote_banding_rules(wb):
    """PLAN.md §22.2 #2: row banding must never outrank a status/priority
    chip or any other meaningful fill. openpyxl numbers rule priority in
    the order rules were added, so a sheet that happened to add its
    banding before its chips (Schedule Data's MeetingsTable did) had every
    chip masked on alternate rows -- the same CF-overrides-fill class as
    the manager build's PIC-color bugs. Rather than relying on every
    sheet's call order, this renumbers every sheet once, at the end of the
    build: all non-banding rules keep their relative order, banding rules
    go last.
    """
    for ws in wb.worksheets:
        rules = [r for cf in ws.conditional_formatting for r in cf.rules]
        for i, r in enumerate(sorted(rules, key=lambda r: (is_banding_rule(r), r.priority)), start=1):
            r.priority = i


def binary_state_rules(column_letter, first_row, last_row, true_label, false_label):
    """§17.4's binary-state poles (good green / muted gray), for fields
    like Active/Inactive or Access Yes/No. `true_label` maps to good,
    `false_label` maps to muted.
    """
    rng = f"{column_letter}{first_row}:{column_letter}{last_row}"
    cell = f"${column_letter}{first_row}"
    return [
        (rng, _chip(f'{cell}="{true_label}"', ds.WORKFLOW_DONE)),
        (rng, _chip(f'{cell}="{false_label}"', ds.WORKFLOW_NOT_STARTED)),
    ]
