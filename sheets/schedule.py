"""Schedule -- the visual work-location/meetings calendar (PLAN.md §13.2).
This module builds the *structure* only (Phase 2): the month grid (reused
from Calendar's mechanism, §8.2), the Work-Location background color
lookup, the meeting-count badge formulas, and the reminder-settings named
cells (§13.5). Click-wiring (Prev/Next/Today, the frmScheduleDay popup,
Application.OnTime reminder scheduling) is VBA, added in Phase 3.

Layout (7 columns, A-G, matching the day grid):
  Row 1 (banner): <Prev | month/year label (merged B1:E1) | Today | Next>
    -- same visual language as Calendar's own banner, but independently
    navigable: this sheet gets its own ScheduleYear/ScheduleMonth named
    cells, distinct from Calendar's SelectedYear/SelectedMonth (two
    separate calendar instances, navigated independently).
  Row 2: weekday headers, Mon-Sun
  Rows 3-14: the 6-week day grid, but **two physical rows per calendar
    week** (confirmed with the user) rather than Calendar's one -- a date
    row (identical in spirit to Calendar's own day cell: value = the live
    DATE formula, format "D", so Location-color CF and any future
    double-click handler key off a clean date value) immediately followed
    by a shorter badge row (a live meeting-count formula, blank when
    zero). Cramming both into one cell's text would leave the cell's
    value a composite string instead of a pure date.
  Row 15: Reminder Settings button (styled only; Phase 3 wires the click)
  Column J (hidden): ScheduleYear/ScheduleMonth + the 3 reminder-setting
    named cells (ReminderFirstOffset/ReminderCount/ReminderInterval,
    §13.5) -- same off-to-the-side hidden-state convention as Calendar's
    own column J and Monthly Media Plan's column P.

**Live-tested design constraint carried over from Campaigns (§10, Phase
2E):** conditional formatting does not retarget relative references per
sub-area in a multi-area (non-contiguous) sqref -- confirmed there via
FormatConditions(1).Formula1. Because the date rows here are no longer
contiguous (a badge row sits between each pair), the adjacent-month
muted-font rule, the Today border, and the Location background color are
each applied **once per week row** (6 separate single-row-contiguous
ranges) rather than one combined range spanning all 6 date rows --
avoiding that bug from the start rather than rediscovering it.
"""

import datetime

from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

from . import design_system as ds
from .week_grid import month_grid_anchor_formula, month_grid_formulas, month_name_full_formula

SHEET_NAME = "Schedule"

_WEEKDAY_HEADERS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
_GRID_FIRST_ROW = 3
_NUM_WEEKS = 6
_ROWS_PER_WEEK = 2  # date row + badge row
_GRID_LAST_ROW = _GRID_FIRST_ROW + _NUM_WEEKS * _ROWS_PER_WEEK - 1  # 14
_WEEKEND_COLS = ("F", "G")
_BUTTON_ROW = _GRID_LAST_ROW + 1  # 15

_STATE_COL = "J"

_DEFAULT_REMINDER_FIRST_OFFSET = 10
_DEFAULT_REMINDER_COUNT = 1
_DEFAULT_REMINDER_INTERVAL = 5

_LOCATION_FILLS = [
    ("Home Office", ds.CATEGORICAL[0]),
    ("Office", ds.CATEGORICAL[1]),
    ("Holiday", ds.CATEGORICAL[2]),
]


def _named(wb, name, sheet_name, cell_ref):
    wb.defined_names[name] = DefinedName(name, attr_text=f"'{sheet_name}'!${cell_ref[0]}${cell_ref[1:]}")


def _date_row(week_index):
    return _GRID_FIRST_ROW + week_index * _ROWS_PER_WEEK


def _badge_row(week_index):
    return _date_row(week_index) + 1


def build(wb):
    """Create the Schedule sheet in `wb` (an openpyxl Workbook). Requires
    Schedule Data to already exist. Returns the created worksheet.
    """
    ws = wb.create_sheet(SHEET_NAME)
    ds.apply_modern_chrome(ws)

    brand_fill = PatternFill("solid", fgColor=ds.BRAND_ACCENT)
    white_bold = Font(name=ds.FONT_NAME, color=ds.BRAND_ACCENT_TEXT, bold=True)
    center = Alignment(horizontal="center", vertical="center")

    # --- Row 1: banner (Prev / month-year label / Today / Next) ---
    ws.row_dimensions[1].height = 26
    for col in range(1, 8):
        ws.cell(row=1, column=col).fill = brand_fill

    prev_cell = ws.cell(row=1, column=1, value=ds.ICON_NAV_PREV)
    prev_cell.font = Font(name=ds.FONT_NAME, color=ds.BRAND_ACCENT_TEXT, bold=True, size=ds.FONT_SIZE_TITLE)
    prev_cell.alignment = center

    ws.merge_cells("B1:E1")
    label_cell = ws.cell(row=1, column=2)
    # Locale-proof month name (PLAN.md §19 Phase 5, found live) -- see
    # week_grid.py's module-level comment for why TEXT(...,"MMMM") isn't
    # safe (renders in the workbook's system locale).
    _month_anchor = "DATE(ScheduleYear,ScheduleMonth,1)"
    label_cell.value = f'={month_name_full_formula(_month_anchor)}&" "&TEXT({_month_anchor},"YYYY")'
    label_cell.font = Font(name=ds.FONT_NAME, color=ds.BRAND_ACCENT_TEXT, bold=True, size=ds.FONT_SIZE_TITLE)
    label_cell.alignment = center

    today_cell = ws.cell(row=1, column=6, value="Today")
    today_cell.font = white_bold
    today_cell.alignment = center

    next_cell = ws.cell(row=1, column=7, value=ds.ICON_NAV_NEXT)
    next_cell.font = Font(name=ds.FONT_NAME, color=ds.BRAND_ACCENT_TEXT, bold=True, size=ds.FONT_SIZE_TITLE)
    next_cell.alignment = center

    # --- Row 2: weekday headers -- divider between columns (§17.11) ---
    for i, label in enumerate(_WEEKDAY_HEADERS):
        cell = ws.cell(row=2, column=i + 1, value=label)
        ds.style_header_cell(cell, wrap=False, divider=(i < len(_WEEKDAY_HEADERS) - 1))

    # --- Rows 3-14: day grid, 2 rows per week ---
    anchor_cell_ref = f"{_STATE_COL}8"  # a scratch cell for the grid anchor formula
    ws[anchor_cell_ref] = month_grid_anchor_formula("ScheduleYear", "ScheduleMonth")
    grid_formulas = month_grid_formulas(anchor_cell_ref)

    day_font = Font(name=ds.FONT_NAME, size=ds.FONT_SIZE_BODY)
    badge_font = Font(name=ds.FONT_NAME, size=ds.FONT_SIZE_CAPTION, color=ds.INK_SECONDARY)
    weekend_fill = PatternFill("solid", fgColor=ds.GRIDLINE)

    work_location_dates = "'Schedule Data'!$A$2:$A$401"
    work_location_values = "'Schedule Data'!$B$2:$B$401"
    meeting_anchor_dates = "'Schedule Data'!$H$2:$H$61"
    meeting_recurrence = "'Schedule Data'!$G$2:$G$61"
    meeting_status = "'Schedule Data'!$I$2:$I$61"

    for week in range(_NUM_WEEKS):
        d_row = _date_row(week)
        b_row = _badge_row(week)
        ws.row_dimensions[d_row].height = 20
        ws.row_dimensions[b_row].height = 14
        row_formulas = grid_formulas[week]

        for c_offset, formula in enumerate(row_formulas):
            col = c_offset + 1
            letter = get_column_letter(col)
            date_cell_ref = f"{letter}{d_row}"

            date_cell = ws.cell(row=d_row, column=col, value=formula)
            date_cell.number_format = "D"
            date_cell.font = day_font
            date_cell.alignment = center
            if letter in _WEEKEND_COLS:
                date_cell.fill = weekend_fill

            one_time_count = f'COUNTIFS({meeting_anchor_dates},{date_cell_ref},{meeting_recurrence},"One-time")'
            recurring_count_raw = (
                f"SUMPRODUCT(({meeting_status}=\"Active\")*(({meeting_recurrence}=\"Daily\")+"
                # "Weekdays" (user-requested, PLAN.md §19 Phase 5): a
                # separate recurrence type from "Daily" (unchanged, still
                # every day), restricted to Monday-Friday --
                # WEEKDAY(date,2) returns 1=Monday..7=Sunday, so <=5
                # covers Mon-Fri. Same matching rule as modSchedule.bas's
                # MeetingOccursOn (the VBA-side equivalent for the day-
                # popup's meeting list and reminders).
                f'(({meeting_recurrence}="Weekdays")*(WEEKDAY({date_cell_ref},2)<=5))+'
                f'(({meeting_recurrence}="Weekly")*(WEEKDAY({meeting_anchor_dates},3)=WEEKDAY({date_cell_ref},3)))+'
                f'(({meeting_recurrence}="Monthly")*(DAY({meeting_anchor_dates})=DAY({date_cell_ref})))))'
            )
            # User-requested addition (PLAN.md §19 Phase 5): a date whose
            # Work Location is Holiday/Sick/Out of Office excludes every
            # RECURRING meeting pattern from counting that day (One-time
            # meetings, in one_time_count above, are unaffected -- an
            # explicitly-picked exact date is a deliberate choice, not a
            # projected pattern). Same three-state check as
            # modSchedule.bas's new IsExcludedDay.
            is_excluded_day = (
                f'(COUNTIFS({work_location_dates},{date_cell_ref},{work_location_values},"Holiday")+'
                f'COUNTIFS({work_location_dates},{date_cell_ref},{work_location_values},"Sick")+'
                f'COUNTIFS({work_location_dates},{date_cell_ref},{work_location_values},"Out of Office"))>0'
            )
            recurring_count = f'IF({is_excluded_day},0,{recurring_count_raw})'
            badge_cell = ws.cell(row=b_row, column=col)
            badge_cell.value = f'=IF(({one_time_count}+{recurring_count})=0,"",{one_time_count}+{recurring_count})'
            badge_cell.font = badge_font
            badge_cell.alignment = center
            # User-requested addition (PLAN.md §19 Phase 5): a small
            # bullet marker ahead of the count, so the badge reads as
            # "there are meetings here" at a glance rather than a bare
            # number that could be mistaken for something else. A number
            # format, not baked into the formula's own text, so it stays
            # purely decorative -- blank (count=0, an empty-string result)
            # shows nothing, matching the cell's existing "invisible when
            # empty" behavior.
            badge_cell.number_format = '"• "0'

        # --- Per-week CF (single contiguous row range each -- see module
        # docstring for why this can't be one combined multi-area range) ---
        week_date_range = f"A{d_row}:G{d_row}"
        top_left = f"A{d_row}"

        ws.conditional_formatting.add(
            week_date_range,
            FormulaRule(formula=[f"MONTH({top_left})<>ScheduleMonth"], font=Font(color=ds.INK_MUTED)),
        )
        today_border = Border(*[Side(style="medium", color=ds.BRAND_ACCENT)] * 4)
        ws.conditional_formatting.add(
            week_date_range, FormulaRule(formula=[f"{top_left}=TODAY()"], border=today_border)
        )
        for label, color in _LOCATION_FILLS:
            location_lookup = f'INDEX({work_location_values},MATCH({top_left},{work_location_dates},0))="{label}"'
            condition = f"AND(COUNTIF({work_location_dates},{top_left})>0,{location_lookup})"
            ws.conditional_formatting.add(
                week_date_range, FormulaRule(formula=[condition], fill=PatternFill("solid", bgColor=color))
            )

    # --- Column widths ---
    for col in range(1, 8):
        ws.column_dimensions[get_column_letter(col)].width = 10

    # --- Row 15: Reminder Settings button (styled only, Phase 3 wires it) ---
    ws.merge_cells(f"A{_BUTTON_ROW}:C{_BUTTON_ROW}")
    settings_button = ws.cell(row=_BUTTON_ROW, column=1, value="Reminder Settings")
    settings_button.font = Font(name=ds.FONT_NAME, size=ds.FONT_SIZE_BODY, bold=True, color=ds.INK_PRIMARY)
    settings_button.fill = PatternFill("solid", fgColor=ds.CARD_SURFACE)
    ds.apply_box_outline(ws, f"A{_BUTTON_ROW}:C{_BUTTON_ROW}", color=ds.INK_MUTED)
    settings_button.alignment = center

    # --- Hidden state (column J): ScheduleYear/ScheduleMonth + reminder settings ---
    today = datetime.date.today()
    ws[f"{_STATE_COL}1"] = today.year
    ws[f"{_STATE_COL}2"] = today.month
    ws[f"{_STATE_COL}3"] = _DEFAULT_REMINDER_FIRST_OFFSET
    ws[f"{_STATE_COL}4"] = _DEFAULT_REMINDER_COUNT
    ws[f"{_STATE_COL}5"] = _DEFAULT_REMINDER_INTERVAL
    ws.column_dimensions[_STATE_COL].hidden = True

    _named(wb, "ScheduleYear", SHEET_NAME, f"{_STATE_COL}1")
    _named(wb, "ScheduleMonth", SHEET_NAME, f"{_STATE_COL}2")
    _named(wb, "ReminderFirstOffset", SHEET_NAME, f"{_STATE_COL}3")
    _named(wb, "ReminderCount", SHEET_NAME, f"{_STATE_COL}4")
    _named(wb, "ReminderInterval", SHEET_NAME, f"{_STATE_COL}5")

    # --- Sheet tab color (core daily-use sheet, §17.9) ---
    ws.sheet_properties.tabColor = ds.TAB_CORE

    return ws
