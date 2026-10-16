"""Daily Log -- "what did I do each day", one section per day of a chosen
week (PLAN.md §22.6, V4, employee only). A read-only VIEW over Archive,
not storage of its own: Archive stays the single flat log every Dashboard
count reads, and this sheet just re-groups it by completion day.

This module builds the fixed chrome only: the week banner (◀ title ▶),
the "This week" and "Copy for stand-up" buttons, and the hidden week
anchor. The body (row 4 down) is rendered by VBA (modDailyLog.RenderWeek)
on every visit and every navigation -- variable-length grouped day
sections with merged headers aren't something legacy-safe formulas can
lay out (no FILTER/SORT in Excel 2016/2019, PLAN.md §1).

Columns: A Time, B Task (as "Parent › Child" for subtasks), C Project /
Category, D Status (✓ done / ⤼ skipped). F (hidden) holds the Monday of
the displayed week, named DailyLogWeekStart.
"""

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName

from . import design_system as ds

SHEET_NAME = "Daily Log"
BODY_FIRST_ROW = 4
WEEK_ANCHOR_CELL = "F1"


def _button(ws, cell_range, label):
    ws.merge_cells(cell_range)
    anchor = ws[cell_range.split(":")[0]]
    anchor.value = label
    anchor.font = Font(name=ds.FONT_NAME, size=ds.FONT_SIZE_BODY, bold=True, color=ds.INK_PRIMARY)
    anchor.fill = PatternFill("solid", fgColor=ds.CARD_SURFACE)
    anchor.alignment = Alignment(horizontal="center", vertical="center")
    ds.apply_box_outline(ws, cell_range, color=ds.INK_MUTED)


def build(wb):
    ws = wb.create_sheet(SHEET_NAME)
    ds.apply_modern_chrome(ws)

    for letter, width in {"A": 10, "B": 56, "C": 22, "D": 14, "E": 3}.items():
        ws.column_dimensions[letter].width = width
    ws.column_dimensions["F"].hidden = True

    # --- Row 1: week banner ◀ [Week 42 · 12.10 – 18.10.2026] ▶ ---
    navy = PatternFill("solid", fgColor=ds.BRAND_ACCENT)
    white_big = Font(name=ds.FONT_NAME, size=ds.FONT_SIZE_TITLE, bold=True, color=ds.BRAND_ACCENT_TEXT)
    center = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 34
    ws.merge_cells("B1:C1")
    for ref, value in (("A1", "◀"), ("B1", "Daily Log"), ("D1", "▶")):
        c = ws[ref]
        c.value = value
        c.font = white_big
        c.alignment = center
    for ref in ("A1", "B1", "C1", "D1"):
        ws[ref].fill = navy

    # --- Row 2: buttons ---
    ws.row_dimensions[2].height = 26
    _button(ws, "A2:B2", "This week")
    _button(ws, "C2:D2", "\U0001F4CB Copy for stand-up")

    # Hidden Monday-of-week anchor; blank = "this week" (VBA fills it on
    # first render and on Workbook_Open).
    ws[WEEK_ANCHOR_CELL].number_format = "DD.MM.YYYY"
    wb.defined_names["DailyLogWeekStart"] = DefinedName(
        "DailyLogWeekStart", attr_text=f"'{SHEET_NAME}'!${WEEK_ANCHOR_CELL[0]}${WEEK_ANCHOR_CELL[1:]}"
    )

    ws.freeze_panes = ws.cell(row=BODY_FIRST_ROW, column=1)

    # Read-only view: nothing here is typed. UserInterfaceOnly is
    # re-asserted on Workbook_Open (doesn't survive a reopen), and
    # RenderWeek unprotects around its own writes anyway.
    ws.protection.sheet = True

    ws.sheet_properties.tabColor = ds.TAB_CORE
    return ws
