"""Time Log -- hidden, internal-only sheet (PLAN.md §2, §6.2 point 4).

Records one row per completed timer session: TaskKey, Date, Duration
(decimal hours). Appended to by Tasks' Timer double-click handler on stop
(§6.3). Read by Dashboard's Time Invested Today/This Week SUMIFS (§4.D,
§6.5) -- a single running Time Spent total per task can't answer "how
much today vs. this week," so each session gets its own dated row here
instead.

Not one of the user's original 13 sheets -- see PLAN.md §2.

A genuine auto-expanding Excel Table (unlike Monthly Media Plan's Gantt or
Campaigns' reserved row blocks, which need rows at fixed positions ahead
of time): starts at the minimum valid size (header + one data row) and
grows via VBA's ListRows.Add as timer sessions are logged, rather than
being pre-allocated to a large fixed size.
"""

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from . import design_system as ds

SHEET_NAME = "Time Log"
HEADERS = ["TaskKey", "Date", "Duration"]


def build(wb):
    """Create the Time Log sheet in `wb` (an openpyxl Workbook), Very
    Hidden per the confirmed decision (accessible only via VBA/the VBA
    editor, not the normal Unhide dialog). Returns the created worksheet.
    """
    ws = wb.create_sheet(SHEET_NAME)

    header_font = Font(name=ds.FONT_NAME, size=ds.FONT_SIZE_HEADER, bold=True, color=ds.BRAND_ACCENT_TEXT)
    header_fill = PatternFill("solid", fgColor=ds.BRAND_ACCENT)
    header_align = Alignment(horizontal="center", vertical="center")

    for col_idx, label in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    # Minimum valid Table size: header row + exactly one data row. VBA
    # grows this via ListRows.Add as timer sessions are logged (Phase 3),
    # which is the idiomatic way to append to a ListObject and have it
    # auto-expand -- pre-allocating many blank rows here would defeat
    # that, unlike the fixed-position cases elsewhere in the workbook.
    ws.cell(row=2, column=2).number_format = "DD.MM.YYYY"
    ws.cell(row=2, column=3).number_format = "0.00"

    ws.column_dimensions[get_column_letter(1)].width = 12
    ws.column_dimensions[get_column_letter(2)].width = 12
    ws.column_dimensions[get_column_letter(3)].width = 10

    table = Table(displayName="TimeLogTable", ref="A1:C2")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)

    ws.sheet_state = "veryHidden"

    return ws
